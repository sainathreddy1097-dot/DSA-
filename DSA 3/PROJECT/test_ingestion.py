from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.services.ingestion import IngestionError, extract_text, segment_clauses, validate_upload


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("1. Scope", level=1)
    document.add_paragraph("Services apply to the customer environment.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Control"
    table.cell(0, 1).text = "Required"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def text_pdf_bytes(text: str) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def test_rejects_oversized_and_unsupported_uploads():
    with pytest.raises(IngestionError, match="10 MiB"):
        validate_upload("large.txt", b"x" * (10 * 1024 * 1024 + 1))
    with pytest.raises(IngestionError, match="PDF, DOCX, or TXT"):
        validate_upload("payload.exe", b"MZ")


def test_filename_never_controls_storage_path():
    document = validate_upload("../../escape.txt", b"Section 1. Scope\nText")

    assert document.filename == "escape.txt"
    assert "/" not in document.filename
    assert "\\" not in document.filename


def test_extracts_utf8_docx_and_text_pdf_content():
    text_document = validate_upload("agreement.txt", "1. Scope\nCafé services".encode())
    word_document = validate_upload("agreement.docx", docx_bytes())
    pdf_document = validate_upload("agreement.pdf", text_pdf_bytes("1. Scope Services apply"))

    assert "Café services" in extract_text(text_document)
    assert "Services apply" in extract_text(word_document)
    assert "Control" in extract_text(word_document)
    assert "Scope Services apply" in extract_text(pdf_document)


def test_rejects_empty_or_image_only_documents():
    empty_pdf = validate_upload("scan.pdf", text_pdf_bytes(""))

    with pytest.raises(IngestionError) as error:
        extract_text(empty_pdf)

    assert error.value.code == "document_text_unavailable"


def test_segments_numbered_legal_headings():
    segments = segment_clauses(
        "1. Scope\nServices apply.\n2. Liability\nLiability is capped."
    )

    assert [(segment.key, segment.title) for segment in segments] == [
        ("C01", "Scope"),
        ("C02", "Liability"),
    ]
    assert segments[1].text == "Liability is capped."


def test_preserves_preamble_and_falls_back_for_unstructured_text():
    with_preamble = segment_clauses("Background terms.\nSECTION 1: Scope\nServices apply.")
    unstructured = segment_clauses("A paragraph without a reliable heading.")

    assert with_preamble[0].title == "Preamble"
    assert with_preamble[1].title == "Scope"
    assert len(unstructured) == 1
    assert unstructured[0].title == "Full document text"
